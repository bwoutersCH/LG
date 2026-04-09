"""Color constants and style helpers for the workbook generator."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── Color hex codes (without #) ─────────────────────────────────────────────
DARK_BLUE = "243347"
ORANGE = "EE742B"
BLUE_GREEN = "239A98"
YELLOW = "FDC400"
SOFT_GREEN = "A8D0C0"
RED = "D9415C"
BLUSH = "F5B0A3"
WHITE = "FFFFFF"
GREY = "C5C6C6"
ANTHRACITE = "575756"
LIGHT_BLUSH = "FDE8E3"

# ── Reusable style objects ───────────────────────────────────────────────────
HEADER_FILL = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

ORANGE_FILL = PatternFill(start_color=ORANGE, end_color=ORANGE, fill_type="solid")
BUTTON_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)

BLUE_GREEN_FILL = PatternFill(start_color=BLUE_GREEN, end_color=BLUE_GREEN, fill_type="solid")
YELLOW_FILL = PatternFill(start_color=YELLOW, end_color=YELLOW, fill_type="solid")
SOFT_GREEN_FILL = PatternFill(start_color=SOFT_GREEN, end_color=SOFT_GREEN, fill_type="solid")
RED_FILL = PatternFill(start_color=RED, end_color=RED, fill_type="solid")
BLUSH_FILL = PatternFill(start_color=BLUSH, end_color=BLUSH, fill_type="solid")
LIGHT_BLUSH_FILL = PatternFill(start_color=LIGHT_BLUSH, end_color=LIGHT_BLUSH, fill_type="solid")
WHITE_FILL = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
GREY_FILL = PatternFill(start_color=GREY, end_color=GREY, fill_type="solid")

BODY_FONT = Font(name="Calibri", size=11, color=ANTHRACITE)
LABEL_FONT = Font(name="Calibri", size=11, color=GREY, bold=True)
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color=DARK_BLUE)
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color=BLUE_GREEN)
ERROR_FONT = Font(name="Calibri", size=11, bold=True, color=RED)
SUCCESS_FONT = Font(name="Calibri", size=11, bold=True, color=BLUE_GREEN)

THIN_BORDER = Border(
    left=Side(style="thin", color=GREY),
    right=Side(style="thin", color=GREY),
    top=Side(style="thin", color=GREY),
    bottom=Side(style="thin", color=GREY),
)


def style_header_row(ws, row, max_col):
    """Apply header styling to a row."""
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def style_button_cell(ws, row, col, text):
    """Style a cell to look like a button (orange fill, white bold text)."""
    cell = ws.cell(row=row, column=col, value=text)
    cell.fill = ORANGE_FILL
    cell.font = BUTTON_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = THIN_BORDER
    return cell


def style_label_cell(ws, row, col, text):
    """Style a cell as a label."""
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = Font(name="Calibri", size=11, bold=True, color=ANTHRACITE)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    return cell


def style_input_cell(ws, row, col):
    """Style a cell as an input field."""
    cell = ws.cell(row=row, column=col)
    cell.border = Border(
        left=Side(style="thin", color=ANTHRACITE),
        right=Side(style="thin", color=ANTHRACITE),
        top=Side(style="thin", color=ANTHRACITE),
        bottom=Side(style="thin", color=ANTHRACITE),
    )
    cell.fill = WHITE_FILL
    cell.font = BODY_FONT
    return cell


def apply_alternating_rows(ws, start_row, end_row, max_col):
    """Apply alternating white/light-blush fills to data rows."""
    for r in range(start_row, end_row + 1):
        fill = LIGHT_BLUSH_FILL if (r - start_row) % 2 == 1 else WHITE_FILL
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).fill = fill
            ws.cell(row=r, column=c).font = BODY_FONT
            ws.cell(row=r, column=c).border = THIN_BORDER
