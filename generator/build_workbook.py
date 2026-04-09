#!/usr/bin/env python3
"""
build_workbook.py — Generates the Time Tracking Excel workbook (.xlsx).
Run: python3 generator/build_workbook.py
Output: TimeTracking.xlsx
"""
import os, sys
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.protection import SheetProtection

sys.path.insert(0, os.path.dirname(__file__))
from colors import (
    HEADER_FILL, HEADER_FONT, HEADER_ALIGN, THIN_BORDER,
    ORANGE_FILL, BUTTON_FONT, BLUE_GREEN_FILL, YELLOW_FILL,
    WHITE_FILL, LIGHT_BLUSH_FILL, BODY_FONT, LABEL_FONT,
    TITLE_FONT, SECTION_FONT, GREY_FILL, RED_FILL,
    style_header_row, style_button_cell, style_label_cell,
    style_input_cell, apply_alternating_rows, ANTHRACITE,
    DARK_BLUE, WHITE, BLUE_GREEN, Alignment, Font, PatternFill
)
from sample_data import (
    TEAMS, TEAM_HEADERS, ACTIVITIES, ACTIVITY_HEADERS, ACTIVITY_NAMES,
    PROJECTS, PROJECT_HEADERS, PROJECT_NAMES, USERS, USER_HEADERS,
    TIME_ENTRY_HEADERS, SESSION_HEADERS, generate_time_entries, TEAM_MAP
)

OUTPUT_FILE = os.path.join(os.path.dirname(os.path.dirname(__file__)), "TimeTracking.xlsx")
SHEET_PASSWORD = "TT#2026!Sec"


def create_table(ws, name, ref, style="TableStyleMedium2"):
    tbl = Table(displayName=name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name=style, showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(tbl)
    return tbl


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_headers(ws, row, headers):
    for c, h in enumerate(headers, 1):
        ws.cell(row=row, column=c, value=h)
    style_header_row(ws, row, len(headers))


def write_data_rows(ws, start_row, data, max_col):
    for r_idx, row_data in enumerate(data):
        row_num = start_row + r_idx
        for c_idx, val in enumerate(row_data):
            cell = ws.cell(row=row_num, column=c_idx + 1, value=val)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            fill = LIGHT_BLUSH_FILL if r_idx % 2 == 1 else WHITE_FILL
            cell.fill = fill
    return start_row + len(data) - 1


# ═══════════════════════════════════════════════════════════════════════════════
# SHEET BUILDERS
# ═══════════════════════════════════════════════════════════════════════════════

def build_login(wb):
    ws = wb.active
    ws.title = "LOGIN"
    set_col_widths(ws, [5, 30, 5, 40])
    ws.sheet_properties.tabColor = DARK_BLUE

    # Title
    ws.merge_cells("A1:D1")
    c = ws.cell(row=1, column=1, value="TIME TRACKING — LOGIN")
    c.font = Font(name="Calibri", size=18, bold=True, color=WHITE)
    c.fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
    c.alignment = Alignment(horizontal="center", vertical="center")

    # Labels + input cells
    style_label_cell(ws, 3, 1, "Username:")
    style_input_cell(ws, 3, 2)
    style_label_cell(ws, 4, 1, "Password:")
    inp = style_input_cell(ws, 4, 2)
    # White text on white bg as visual password mask
    inp.font = Font(name="Calibri", size=11, color=WHITE)

    # Button placeholder
    style_button_cell(ws, 5, 2, "[ Log In ]  ->  Run Login.ts > runLogin()")
    style_button_cell(ws, 7, 2, "[ Log Out ]  ->  Run Login.ts > runLogout()")

    # Error message cell
    ws.cell(row=6, column=2, value="").font = Font(name="Calibri", size=11, color="D9415C")

    # Instructions
    ws.cell(row=9, column=1, value="Default credentials (all users):").font = SECTION_FONT
    ws.cell(row=10, column=1, value="Password: Pass1234").font = BODY_FONT
    ws.cell(row=11, column=1, value="Usernames: alice.admin, bob.director, dave.manager, hank.teamlead, leo.employee, etc.").font = BODY_FONT
    ws.cell(row=12, column=1, value="See README sheet for full user list.").font = BODY_FONT
    return ws


def build_session(wb):
    ws = wb.create_sheet("SESSION")
    write_headers(ws, 1, SESSION_HEADERS)
    set_col_widths(ws, [12, 18, 12, 12, 25])
    ref = f"A1:{get_column_letter(len(SESSION_HEADERS))}2"
    # Add an empty data row so table has a body
    for c in range(1, len(SESSION_HEADERS) + 1):
        ws.cell(row=2, column=c, value="")
    create_table(ws, "SessionTable", ref)
    ws.protection = SheetProtection(sheet=True, password=SHEET_PASSWORD)
    ws.sheet_state = "hidden"
    return ws


def build_users_db(wb):
    ws = wb.create_sheet("USERS_DB")
    write_headers(ws, 1, USER_HEADERS)
    set_col_widths(ws, [10, 22, 22, 20, 12, 12, 10, 8])
    end_row = write_data_rows(ws, 2, USERS, len(USER_HEADERS))
    ref = f"A1:{get_column_letter(len(USER_HEADERS))}{end_row}"
    create_table(ws, "UsersTable", ref)
    ws.protection = SheetProtection(sheet=True, password=SHEET_PASSWORD)
    ws.sheet_state = "hidden"
    return ws


def build_projects_db(wb):
    ws = wb.create_sheet("PROJECTS_DB")
    write_headers(ws, 1, PROJECT_HEADERS)
    set_col_widths(ws, [12, 22, 10, 8])
    end_row = write_data_rows(ws, 2, PROJECTS, len(PROJECT_HEADERS))
    ref = f"A1:{get_column_letter(len(PROJECT_HEADERS))}{end_row}"
    create_table(ws, "ProjectsTable", ref)
    ws.protection = SheetProtection(sheet=True, password=SHEET_PASSWORD)
    ws.sheet_state = "hidden"
    return ws


def build_activities_db(wb):
    ws = wb.create_sheet("ACTIVITIES_DB")
    write_headers(ws, 1, ACTIVITY_HEADERS)
    set_col_widths(ws, [12, 22, 8])
    end_row = write_data_rows(ws, 2, ACTIVITIES, len(ACTIVITY_HEADERS))
    ref = f"A1:{get_column_letter(len(ACTIVITY_HEADERS))}{end_row}"
    create_table(ws, "ActivitiesTable", ref)
    ws.protection = SheetProtection(sheet=True, password=SHEET_PASSWORD)
    ws.sheet_state = "hidden"
    return ws


def build_teams_db(wb):
    ws = wb.create_sheet("TEAMS_DB")
    write_headers(ws, 1, TEAM_HEADERS)
    set_col_widths(ws, [10, 18, 14, 14])
    end_row = write_data_rows(ws, 2, TEAMS, len(TEAM_HEADERS))
    ref = f"A1:{get_column_letter(len(TEAM_HEADERS))}{end_row}"
    create_table(ws, "TeamsTable", ref)
    ws.protection = SheetProtection(sheet=True, password=SHEET_PASSWORD)
    ws.sheet_state = "hidden"
    return ws


def build_time_entry(wb):
    ws = wb.create_sheet("TIME_ENTRY")
    ws.sheet_properties.tabColor = BLUE_GREEN
    set_col_widths(ws, [16, 18, 14, 22, 16, 20, 18, 8, 30, 22, 16, 22])

    # ── Input area ───────────────────────────────────────────────────────────
    ws.merge_cells("A1:L1")
    c = ws.cell(row=1, column=1, value="TIME ENTRY")
    c.font = TITLE_FONT
    c.fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
    c.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    c.alignment = Alignment(horizontal="center")

    labels = ["Date:", "Project:", "Activity:", "Hours (1-16):", "Notes:"]
    for i, lbl in enumerate(labels):
        style_label_cell(ws, 2 + i, 1, lbl)
        style_input_cell(ws, 2 + i, 2)

    # Set today's date as default
    ws.cell(row=2, column=2, value=date.today().isoformat())

    # Data validation for Project dropdown
    proj_list = ",".join(PROJECT_NAMES)
    dv_proj = DataValidation(type="list", formula1=f'"{proj_list}"', allow_blank=False)
    dv_proj.error = "Select a project from the list"
    dv_proj.errorTitle = "Invalid Project"
    ws.add_data_validation(dv_proj)
    dv_proj.add(ws["B3"])

    # Data validation for Activity dropdown
    act_list = ",".join(ACTIVITY_NAMES)
    dv_act = DataValidation(type="list", formula1=f'"{act_list}"', allow_blank=False)
    dv_act.error = "Select an activity from the list"
    dv_act.errorTitle = "Invalid Activity"
    ws.add_data_validation(dv_act)
    dv_act.add(ws["B4"])

    # Data validation for Hours
    dv_hrs = DataValidation(type="whole", operator="between", formula1="1", formula2="16")
    dv_hrs.error = "Hours must be 1-16"
    dv_hrs.errorTitle = "Invalid Hours"
    ws.add_data_validation(dv_hrs)
    dv_hrs.add(ws["B5"])

    # Buttons
    style_button_cell(ws, 7, 2, "[ Add Entry ]  ->  TimeEntry.ts > addEntry()")
    style_button_cell(ws, 8, 2, "[ My Entries This Week ]  ->  TimeEntry.ts > filterMyEntries()")
    style_button_cell(ws, 9, 2, "[ Clear Filters ]  ->  TimeEntry.ts > clearFilters()")

    # Delete entry input
    style_label_cell(ws, 2, 5, "Delete EntryID:")
    style_input_cell(ws, 2, 6)
    style_button_cell(ws, 3, 5, "[ Delete Entry ]  ->  TimeEntry.ts > deleteEntry()")

    # Status message cell
    ws.cell(row=7, column=4, value="").font = BODY_FONT

    # ── Data table ───────────────────────────────────────────────────────────
    ws.cell(row=11, column=1, value="TIME ENTRIES").font = SECTION_FONT
    write_headers(ws, 12, TIME_ENTRY_HEADERS)

    entries = generate_time_entries()
    end_row = write_data_rows(ws, 13, entries, len(TIME_ENTRY_HEADERS))
    ref = f"A12:{get_column_letter(len(TIME_ENTRY_HEADERS))}{end_row}"
    create_table(ws, "TimeEntryTable", ref)

    ws.sheet_state = "hidden"
    return ws


def build_daily_check(wb):
    ws = wb.create_sheet("DAILY_CHECK")
    ws.sheet_properties.tabColor = BLUE_GREEN
    set_col_widths(ws, [18, 16, 14])

    ws.merge_cells("A1:C1")
    c = ws.cell(row=1, column=1, value="DAILY CHECK — Weekly Hours")
    c.font = Font(name="Calibri", size=14, bold=True, color=WHITE)
    c.fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
    c.alignment = Alignment(horizontal="center")

    headers = ["Day", "Date", "Hours"]
    write_headers(ws, 2, headers)

    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    today = date.today()
    monday = today - timedelta(days=today.weekday())
    for i, day_name in enumerate(days):
        row = 3 + i
        d = monday + timedelta(days=i)
        ws.cell(row=row, column=1, value=day_name).font = BODY_FONT
        ws.cell(row=row, column=2, value=d.isoformat()).font = BODY_FONT
        ws.cell(row=row, column=3, value=0).font = BODY_FONT
        for cc in range(1, 4):
            ws.cell(row=row, column=cc).border = THIN_BORDER

    ws.cell(row=8, column=1, value="TOTAL").font = Font(name="Calibri", size=11, bold=True, color=ANTHRACITE)
    ws.cell(row=8, column=3, value=0).font = Font(name="Calibri", size=11, bold=True, color=ANTHRACITE)

    style_button_cell(ws, 9, 2, "[ Validate Week ]  ->  TimeEntry.ts > validateWeek()")

    ws.cell(row=10, column=1, value="").font = BODY_FONT  # validation result cell

    ws.sheet_state = "hidden"
    return ws


def build_manager_view(wb):
    ws = wb.create_sheet("MANAGER_VIEW")
    ws.sheet_properties.tabColor = BLUE_GREEN
    set_col_widths(ws, [16, 18, 14, 22, 16, 20, 18, 8, 30, 22, 16, 22])

    ws.merge_cells("A1:L1")
    c = ws.cell(row=1, column=1, value="MANAGER VIEW — Team Time Entries")
    c.font = Font(name="Calibri", size=14, bold=True, color=WHITE)
    c.fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
    c.alignment = Alignment(horizontal="center")

    # Filter controls
    style_label_cell(ws, 2, 1, "Team:")
    style_input_cell(ws, 2, 2)
    team_names = ["All"] + [t[1] for t in TEAMS]
    dv_team = DataValidation(type="list", formula1=f'"{",".join(team_names)}"', allow_blank=True)
    ws.add_data_validation(dv_team)
    dv_team.add(ws["B2"])

    style_label_cell(ws, 2, 3, "Employee:")
    style_input_cell(ws, 2, 4)

    style_label_cell(ws, 2, 5, "From:")
    style_input_cell(ws, 2, 6)
    style_label_cell(ws, 2, 7, "To:")
    style_input_cell(ws, 2, 8)

    style_button_cell(ws, 3, 2, "[ Apply Filter ]  ->  Reports.ts > filterManagerView()")

    ws.cell(row=4, column=1, value="").font = BODY_FONT  # status message

    # Data table
    write_headers(ws, 6, TIME_ENTRY_HEADERS)
    # Empty row for table body
    for cc in range(1, len(TIME_ENTRY_HEADERS) + 1):
        ws.cell(row=7, column=cc, value="").font = BODY_FONT
    ref = f"A6:{get_column_letter(len(TIME_ENTRY_HEADERS))}7"
    create_table(ws, "ManagerViewTable", ref)

    # Correction section
    ws.merge_cells("A12:F12")
    c = ws.cell(row=12, column=1, value="CORRECT ENTRY")
    c.font = SECTION_FONT
    c.fill = BLUE_GREEN_FILL
    c.font = Font(name="Calibri", size=12, bold=True, color=WHITE)

    style_label_cell(ws, 13, 1, "Entry ID:")
    style_label_cell(ws, 13, 3, "New Hours:")
    style_label_cell(ws, 13, 5, "Reason:")

    style_input_cell(ws, 14, 2)
    style_input_cell(ws, 14, 4)
    style_input_cell(ws, 14, 6)

    style_button_cell(ws, 15, 2, "[ Save Correction ]  ->  Reports.ts > correctEntry()")
    ws.cell(row=16, column=2, value="").font = BODY_FONT  # message cell

    ws.sheet_state = "hidden"
    return ws


def build_reports(wb):
    ws = wb.create_sheet("REPORTS")
    ws.sheet_properties.tabColor = BLUE_GREEN
    set_col_widths(ws, [22, 30, 18, 18, 18, 18, 18, 18])

    ws.merge_cells("A1:H1")
    c = ws.cell(row=1, column=1, value="REPORTS")
    c.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    c.fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
    c.alignment = Alignment(horizontal="center")

    # Period selector
    style_label_cell(ws, 2, 1, "Period:")
    style_input_cell(ws, 2, 2)
    periods = "This week,This month,This quarter,YTD,Full year"
    dv_period = DataValidation(type="list", formula1=f'"{periods}"', allow_blank=False)
    ws.add_data_validation(dv_period)
    dv_period.add(ws["B2"])
    ws["B2"].value = "This week"

    # Breakdown selector
    style_label_cell(ws, 3, 1, "Breakdown:")
    style_input_cell(ws, 3, 2)
    breakdowns = "By Employee,By Project,By Activity,By Team"
    dv_break = DataValidation(type="list", formula1=f'"{breakdowns}"', allow_blank=False)
    ws.add_data_validation(dv_break)
    dv_break.add(ws["B3"])
    ws["B3"].value = "By Employee"

    # Team filter
    style_label_cell(ws, 4, 1, "Team:")
    style_input_cell(ws, 4, 2)
    team_names = ["All"] + [t[1] for t in TEAMS]
    dv_team = DataValidation(type="list", formula1=f'"{",".join(team_names)}"', allow_blank=True)
    ws.add_data_validation(dv_team)
    dv_team.add(ws["B4"])
    ws["B4"].value = "All"

    style_button_cell(ws, 5, 2, "[ Generate Report ]  ->  Reports.ts > generateReport()")

    ws.cell(row=7, column=1, value="Report output appears below after running the script.").font = LABEL_FONT

    ws.sheet_state = "hidden"
    return ws


def build_admin(wb):
    ws = wb.create_sheet("ADMIN")
    ws.sheet_properties.tabColor = "EE742B"
    set_col_widths(ws, [18, 30, 5, 18, 5, 40])

    ws.merge_cells("A1:F1")
    c = ws.cell(row=1, column=1, value="ADMIN PANEL")
    c.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    c.fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
    c.alignment = Alignment(horizontal="center")

    # ── User management section ──────────────────────────────────────────────
    ws.cell(row=2, column=1, value="MANAGE USERS").font = SECTION_FONT
    labels_u = ["UserID (blank=new):", "Full Name:", "Username:", "Password:", "Role:", "ManagerID:", "TeamID:", "Active (TRUE/FALSE):"]
    for i, lbl in enumerate(labels_u):
        style_label_cell(ws, 3 + i, 1, lbl)
        style_input_cell(ws, 3 + i, 2)

    # Role dropdown
    dv_role = DataValidation(type="list", formula1='"Admin,Director,Manager,TeamLead,Employee"', allow_blank=False)
    ws.add_data_validation(dv_role)
    dv_role.add(ws["B7"])

    style_button_cell(ws, 11, 2, "[ Save User ]  ->  Admin.ts > saveUser()")
    ws.cell(row=3, column=6, value="").font = BODY_FONT  # message cell

    # ── Project management section ───────────────────────────────────────────
    ws.cell(row=12, column=1, value="MANAGE PROJECTS").font = SECTION_FONT
    labels_p = ["ProjectID (blank=new):", "Project Name:", "TeamID:", "Active (TRUE/FALSE):"]
    for i, lbl in enumerate(labels_p):
        style_label_cell(ws, 13 + i, 1, lbl)
        style_input_cell(ws, 13 + i, 2)

    style_button_cell(ws, 17, 2, "[ Save Project ]  ->  Admin.ts > saveProject()")
    ws.cell(row=13, column=6, value="").font = BODY_FONT

    # ── Activity management section ──────────────────────────────────────────
    ws.cell(row=18, column=1, value="MANAGE ACTIVITIES").font = SECTION_FONT
    labels_a = ["ActivityID (blank=new):", "Activity Name:", "Active (TRUE/FALSE):"]
    for i, lbl in enumerate(labels_a):
        style_label_cell(ws, 18 + i + 1, 1, lbl)  # rows 19, 20, 21
        style_input_cell(ws, 18 + i + 1, 2)

    style_button_cell(ws, 22, 2, "[ Save Activity ]  ->  Admin.ts > saveActivity()")
    ws.cell(row=19, column=6, value="").font = BODY_FONT

    # ── Utility buttons ──────────────────────────────────────────────────────
    ws.cell(row=24, column=1, value="UTILITIES").font = SECTION_FONT
    style_button_cell(ws, 25, 2, "[ Reset Session ]  ->  Admin.ts > resetSession()")
    style_button_cell(ws, 26, 2, "[ Export All Data ]  ->  Admin.ts > exportAllData()")
    style_button_cell(ws, 27, 2, "[ Initialize Sample Data ]  ->  InitializeWorkbook.ts > initializeWorkbook()")
    ws.cell(row=22, column=6, value="").font = BODY_FONT

    ws.sheet_state = "hidden"
    return ws


def build_readme(wb):
    ws = wb.create_sheet("README")
    ws.sheet_properties.tabColor = "239A98"
    set_col_widths(ws, [4, 90])

    ws.merge_cells("A1:B1")
    c = ws.cell(row=1, column=1, value="TIME TRACKING WORKBOOK — README")
    c.font = Font(name="Calibri", size=18, bold=True, color=WHITE)
    c.fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
    c.alignment = Alignment(horizontal="center")

    content = [
        "",
        "FIRST-RUN SETUP",
        "1. Upload this .xlsx file to a SharePoint document library.",
        "2. Open in Excel for the Web.",
        "3. Go to Automate tab > New Script. Create each script from the /scripts folder:",
        "   - Utils.ts, Login.ts, TimeEntry.ts, Reports.ts, Admin.ts, InitializeWorkbook.ts",
        "4. Run InitializeWorkbook.ts > initializeWorkbook() ONCE to verify sample data is loaded.",
        "   (Sample data is pre-loaded in this file; re-running will create duplicates.)",
        "5. Assign each button to its script function (right-click button cell > Assign Script).",
        "6. Log in with: alice.admin / Pass1234 (Admin) to verify everything works.",
        "",
        "DEFAULT CREDENTIALS",
        "All sample users have password: Pass1234",
        "",
        "SAMPLE USERNAMES",
        "Admin:      alice.admin",
        "Directors:  bob.director, carol.director",
        "Managers:   dave.manager, eve.manager, frank.manager, grace.manager",
        "TeamLeads:  hank.teamlead, ivy.teamlead, jack.teamlead, kate.teamlead",
        "Employees:  leo.employee, mia.employee, noah.employee, olivia.employee,",
        "            pete.employee, quinn.employee, ryan.employee, sara.employee,",
        "            tom.employee, uma.employee, vic.employee, wendy.employee,",
        "            xander.employee, yara.employee",
        "",
        "HOW TO ADD/MODIFY USERS",
        "1. Log in as Admin.",
        "2. Go to the ADMIN sheet.",
        "3. Fill in user fields. Leave UserID blank for a new user.",
        "4. Click [ Save User ] button (runs Admin.ts > saveUser()).",
        "5. To deactivate a user, set Active to FALSE and save.",
        "",
        "HOW TO ASSIGN MANAGERS",
        "- Each user has a ManagerID field pointing to another user's UserID.",
        "- Set ManagerID = U002 to make Bob Director the manager.",
        "- The hierarchy is resolved dynamically — any depth is supported.",
        "- Admins (ManagerID blank) are at the top of the hierarchy.",
        "",
        "SCRIPT DESCRIPTIONS",
        "Utils.ts          Shared utilities: hash, session, hierarchy, date helpers",
        "Login.ts          Authentication: runLogin(), runLogout(), refreshDailyCheck()",
        "TimeEntry.ts      Time logging: addEntry(), filterMyEntries(), validateWeek(), deleteEntry()",
        "Reports.ts        Manager tools: filterManagerView(), correctEntry(), generateReport()",
        "Admin.ts          Admin ops: saveUser(), saveProject(), saveActivity(), resetSession(), exportAllData()",
        "InitializeWorkbook.ts  One-time setup: populates all DB tables with sample data",
        "",
        "SHEET ACCESS BY ROLE",
        "LOGIN:         Always visible (hidden after login)",
        "README:        Always visible",
        "TIME_ENTRY:    All logged-in users",
        "DAILY_CHECK:   All logged-in users",
        "MANAGER_VIEW:  TeamLead, Manager, Director, Admin",
        "REPORTS:       Manager, Director, Admin",
        "ADMIN:         Admin only",
        "DB sheets:     Hidden, protected (Admin can unhide)",
        "SESSION:       Hidden, protected (system use only)",
        "",
        "HIERARCHY RULES",
        "- Managers see time entries for all users in their subtree (recursive).",
        "- Admins see and edit everything.",
        "- Employees see only their own entries.",
        "- The hierarchy depth is not hardcoded — it walks ManagerID chains.",
        "",
        "VALIDATION RULES",
        "- Hours per entry: 1-16 (whole numbers only)",
        "- Date: within 30 days of today",
        "- Days with < 8h logged are highlighted yellow",
        "- Duplicate date+project+activity entries trigger a warning",
        "",
        "COLOR LEGEND",
        "Dark blue (#243347):  Headers and title bars",
        "Orange (#EE742B):     Buttons / action triggers",
        "Blue-green (#239A98): Section dividers, success messages",
        "Yellow (#FDC400):     Warnings (< 8h days)",
        "Red (#D9415C):        Errors and flags",
        "Soft green (#A8D0C0): Approved / on-track rows",
        "",
        "KNOWN LIMITATIONS",
        "- Office Scripts cannot intercept workbook close events; use the Log Out button.",
        "- Password masking is visual only (white text on white); not true password fields.",
        "- The hash function is deterministic but NOT cryptographic-grade.",
        "- Office Scripts run sequentially; no real-time collaboration locking.",
        "- Button cells are placeholders; assign scripts via Automate > Assign Script.",
        "- Data validation dropdowns are static; run a refresh script after adding projects/activities.",
        "- Sheet protection uses a simple password; it prevents accidental edits, not determined attacks.",
    ]

    for i, line in enumerate(content):
        row = 2 + i
        cell = ws.cell(row=row, column=2, value=line)
        if line and line == line.upper() and not line.startswith(" ") and not line.startswith("-"):
            cell.font = SECTION_FONT
        elif line.startswith("   "):
            cell.font = BODY_FONT
        else:
            cell.font = BODY_FONT

    ws.protection = SheetProtection(sheet=True, password=SHEET_PASSWORD)
    return ws


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════════

def main():
    print("Building Time Tracking workbook...")
    wb = Workbook()

    # Build all sheets in order
    build_login(wb)
    build_session(wb)
    build_users_db(wb)
    build_projects_db(wb)
    build_activities_db(wb)
    build_teams_db(wb)
    build_time_entry(wb)
    build_daily_check(wb)
    build_manager_view(wb)
    build_reports(wb)
    build_admin(wb)
    build_readme(wb)

    # Save
    wb.save(OUTPUT_FILE)
    print(f"Workbook saved to: {OUTPUT_FILE}")
    print(f"Sheets: {[ws.title for ws in wb.worksheets]}")

    # Count time entries
    te_ws = wb["TIME_ENTRY"]
    data_rows = te_ws.max_row - 12  # subtract header area
    print(f"Time entries generated: {data_rows}")
    print("Done!")


if __name__ == "__main__":
    main()
